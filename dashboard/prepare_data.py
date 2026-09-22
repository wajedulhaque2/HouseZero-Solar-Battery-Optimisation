"""Build the compact hourly snapshot from the repository's cached Excel model."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "workbook" / "HouseZero_Solar_Battery_Optimisation_Portfolio.xlsx"
OUTPUT = Path(__file__).resolve().parent / "data" / "hourly.csv.gz"


def extract(path: Path = WORKBOOK) -> pd.DataFrame:
    sheet = load_workbook(path, read_only=True, data_only=True)["Hourly_Data"]
    rows = sheet.iter_rows(values_only=True)
    for _ in range(3):
        next(rows)
    records = [(r[1], r[6], r[7]) for r in rows if r[1] is not None]
    frame = pd.DataFrame(records, columns=["local_time", "demand_kwh", "pv_kwh"])
    frame["local_time"] = pd.to_datetime(frame["local_time"])
    for field in ("demand_kwh", "pv_kwh"):
        frame[field] = pd.to_numeric(frame[field], errors="raise")
    if len(frame) != 8784 or frame.local_time.isna().any() or frame[["demand_kwh", "pv_kwh"]].isna().any().any():
        raise ValueError("Cached hourly source has unexpected length or missing values")
    return frame


if __name__ == "__main__":
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    data = extract()
    data.to_csv(OUTPUT, index=False, compression={"method": "gzip", "mtime": 0})
    print(f"{len(data):,} hours; demand {data.demand_kwh.sum():,.2f} kWh; PV {data.pv_kwh.sum():,.2f} kWh")
